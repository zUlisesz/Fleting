"""Small Python API for reading and editing the inventory workbook.

Example::

    from scrip import ExcelAPI

    with ExcelAPI("Inventario.xlsx") as inventory:
        matches = inventory.search("cutter", columns=["Producto"])
        inventory.update_by_code("100098", {"Existencia": 5})
        inventory.save()
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from copy import copy
import os
from pathlib import Path
import tempfile
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook import Workbook


DEFAULT_FILE = Path(__file__).with_name("Inventario.xlsx")
INVENTORY_HEADERS = (
    "Código",
    "Producto",
    "P. Costo",
    "P. Venta",
    "P. Mayoreo",
    "Departamento",
    "Existencia",
    "Inv. Mínimo",
    "Inv. Máximo",
    "Tipo de Venta",
)


class ExcelAPI:
    """Read and modify rows in an XLSX sheet using its header names."""

    def __init__(
        self,
        filepath: str | Path = DEFAULT_FILE,
        sheet_name: str | None = None,
        header_row: int = 1,
        key_column: str = "Código",
        expected_headers: Sequence[str] | None = None,
    ) -> None:
        self.filepath = Path(filepath).expanduser()
        if not self.filepath.is_file():
            raise FileNotFoundError(f"Workbook not found: {self.filepath}")
        if header_row < 1:
            raise ValueError("header_row must be 1 or greater")

        self.workbook: Workbook = load_workbook(self.filepath)
        if sheet_name is None:
            self.sheet = self.workbook.active
        elif sheet_name in self.workbook.sheetnames:
            self.sheet = self.workbook[sheet_name]
        else:
            self.workbook.close()
            raise KeyError(f"Worksheet not found: {sheet_name}")

        self.header_row = header_row
        self.key_column = key_column
        try:
            self.headers = self._read_headers()
            if expected_headers is not None and tuple(self.headers) != tuple(expected_headers):
                raise ValueError(
                    "The worksheet columns do not match the inventory structure. "
                    f"Expected: {', '.join(expected_headers)}. "
                    f"Found: {', '.join(self.headers)}."
                )
        except Exception:
            self.workbook.close()
            raise

    def _read_headers(self) -> list[str]:
        headers = [cell.value for cell in self.sheet[self.header_row]]
        while headers and headers[-1] is None:
            headers.pop()
        if not headers:
            raise ValueError(f"No headers found on row {self.header_row}")
        if any(not isinstance(header, str) or not header.strip() for header in headers):
            raise ValueError("Every column header must be a non-empty string")
        if len(set(headers)) != len(headers):
            raise ValueError("Column headers must be unique")
        if self.key_column not in headers:
            self.key_column = ""
        return headers

    def _validate_fields(self, values: Mapping[str, Any]) -> None:
        unknown = set(values) - set(self.headers)
        if unknown:
            raise KeyError(f"Unknown column(s): {', '.join(sorted(unknown))}")

    def _row_record(self, row_number: int) -> dict[str, Any]:
        return {
            header: self.sheet.cell(row=row_number, column=column).value
            for column, header in enumerate(self.headers, start=1)
        }

    def _row_is_empty(self, row_number: int) -> bool:
        return all(value is None for value in self._row_record(row_number).values())

    def _matches(self, row_number: int, criteria: Mapping[str, Any]) -> bool:
        record = self._row_record(row_number)
        return all(record[column] == expected for column, expected in criteria.items())

    def get_all(
        self,
        filters: Mapping[str, Any] | None = None,
        *,
        limit: int | None = None,
        offset: int = 0,
    ) -> list[dict[str, Any]]:
        """Return non-empty rows, optionally filtered and paginated."""
        filters = filters or {}
        self._validate_fields(filters)
        if offset < 0 or (limit is not None and limit < 0):
            raise ValueError("limit and offset cannot be negative")

        rows = [
            self._row_record(row_number)
            for row_number in range(self.header_row + 1, self.sheet.max_row + 1)
            if not self._row_is_empty(row_number)
            and self._matches(row_number, filters)
        ]
        return rows[offset:] if limit is None else rows[offset : offset + limit]

    def find(self, criteria: Mapping[str, Any]) -> list[dict[str, Any]]:
        """Return all rows whose column values exactly match ``criteria``."""
        return self.get_all(filters=criteria)

    def get_one(self, criteria: Mapping[str, Any]) -> dict[str, Any] | None:
        """Return one matching row, or ``None``; reject ambiguous matches."""
        matches = self.find(criteria)
        if len(matches) > 1:
            raise ValueError(f"Expected one row for {dict(criteria)!r}; found {len(matches)}")
        return matches[0] if matches else None

    def get_by_code(self, code: Any) -> dict[str, Any] | None:
        """Find one inventory item by its ``Código`` value."""
        if not self.key_column:
            raise ValueError("This worksheet does not have a Código column")
        return self.get_one({self.key_column: code})

    def search(
        self,
        query: str,
        *,
        columns: Sequence[str] | None = None,
        case_sensitive: bool = False,
    ) -> list[dict[str, Any]]:
        """Find rows containing text in one or more columns."""
        search_columns = list(columns) if columns is not None else self.headers
        self._validate_fields(dict.fromkeys(search_columns, None))
        needle = query if case_sensitive else query.casefold()
        results = []
        for record in self.get_all():
            for column in search_columns:
                value = record[column]
                if value is None:
                    continue
                text = str(value) if case_sensitive else str(value).casefold()
                if needle in text:
                    results.append(record)
                    break
        return results

    def add(self, record: Mapping[str, Any]) -> dict[str, Any]:
        """Append a row and return the completed record."""
        if not record:
            raise ValueError("record cannot be empty")
        self._validate_fields(record)

        if self.key_column and record.get(self.key_column) is not None:
            if self.get_by_code(record[self.key_column]) is not None:
                raise ValueError(f"Duplicate {self.key_column}: {record[self.key_column]!r}")

        row_number = max(self.sheet.max_row + 1, self.header_row + 1)
        for column, header in enumerate(self.headers, start=1):
            cell = self.sheet.cell(row=row_number, column=column, value=record.get(header))
            if row_number > self.header_row + 1:
                previous_cell = self.sheet.cell(row=row_number - 1, column=column)
                if previous_cell.has_style:
                    cell.font = copy(previous_cell.font)
                    cell.fill = copy(previous_cell.fill)
                    cell.border = copy(previous_cell.border)
                    cell.alignment = copy(previous_cell.alignment)
                    cell.protection = copy(previous_cell.protection)
                    cell.number_format = previous_cell.number_format
        return self._row_record(row_number)

    def update(self, criteria: Mapping[str, Any], changes: Mapping[str, Any]) -> int:
        """Update every row matching ``criteria`` and return the row count."""
        if not criteria:
            raise ValueError("criteria cannot be empty; refusing to update every row")
        if not changes:
            raise ValueError("changes cannot be empty")
        self._validate_fields(criteria)
        self._validate_fields(changes)

        updated = 0
        for row_number in range(self.header_row + 1, self.sheet.max_row + 1):
            if self._row_is_empty(row_number) or not self._matches(row_number, criteria):
                continue
            for column, value in changes.items():
                column_number = self.headers.index(column) + 1
                self.sheet.cell(row=row_number, column=column_number).value = value
            updated += 1
        return updated

    def update_by_code(self, code: Any, changes: Mapping[str, Any]) -> int:
        """Update an inventory item identified by its ``Código``."""
        if not self.key_column:
            raise ValueError("This worksheet does not have a Código column")
        if self.key_column in changes and changes[self.key_column] != code:
            raise ValueError("Changing Código through update_by_code is not supported")
        return self.update({self.key_column: code}, changes)

    def delete(self, criteria: Mapping[str, Any]) -> int:
        """Delete every row matching ``criteria`` and return the row count."""
        if not criteria:
            raise ValueError("criteria cannot be empty; refusing to delete every row")
        self._validate_fields(criteria)
        matching_rows = [
            row_number
            for row_number in range(self.header_row + 1, self.sheet.max_row + 1)
            if not self._row_is_empty(row_number) and self._matches(row_number, criteria)
        ]
        for row_number in reversed(matching_rows):
            self.sheet.delete_rows(row_number)
        return len(matching_rows)

    def delete_by_code(self, code: Any) -> int:
        """Delete an inventory item identified by its ``Código``."""
        if not self.key_column:
            raise ValueError("This worksheet does not have a Código column")
        return self.delete({self.key_column: code})

    def save(self, filepath: str | Path | None = None) -> Path:
        """Save atomically; a different path becomes the current workbook path."""
        destination = Path(filepath).expanduser() if filepath is not None else self.filepath
        if destination.suffix.lower() != ".xlsx":
            raise ValueError("The destination must be an .xlsx file")
        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary_path = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=f".{destination.stem}.",
                suffix=".xlsx",
                dir=destination.parent,
                delete=False,
            ) as temporary_file:
                temporary_path = Path(temporary_file.name)
            self.workbook.save(temporary_path)
            if destination.exists():
                os.chmod(temporary_path, destination.stat().st_mode & 0o777)
            os.replace(temporary_path, destination)
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
        self.filepath = destination
        return destination

    def close(self) -> None:
        self.workbook.close()

    def __enter__(self) -> ExcelAPI:
        return self

    def __exit__(self, exc_type: Any, exc_value: Any, traceback: Any) -> None:
        self.close()


if __name__ == "__main__":
    with ExcelAPI() as inventory:
        no_department = inventory.find({"Departamento": "- Sin Departamento -"})
        for item in no_department:
            print(
                f"Código: {item['Código']}, Producto: {item['Producto']}, "
                f"Departamento: {item['Departamento']}"
            )
